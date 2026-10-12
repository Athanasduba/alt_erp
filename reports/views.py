from django.shortcuts import render, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.db.models import Sum, Count, F, Q
from django.utils.dateparse import parse_date
from customers.models import Customer
from agents.models import Agent
from datetime import datetime
from django.utils import timezone
from invoices.models import Invoice
from payments.models import Payment
from accounting.models import LedgerEntry, CustomerAccount, CommissionPayout
from motorcycles.models import Motorcycle, MotorcycleHistory
from warehouses.models import Warehouse
# ✅ Import export utilities
from .utils import export_to_excel, export_to_pdf
#from utils.filters import get_date_range, apply_date_filter
from django.utils.formats import number_format
from django.http import HttpResponse
from django.utils.timezone import now
from reports.utils import export_movement_excel, export_movement_pdf
import openpyxl
from django.template.loader import render_to_string
from xhtml2pdf import pisa
from openpyxl.styles import Font, Alignment, PatternFill




# -----------------------------
# HELPER: DATE RANGE FILTER
# -----------------------------
def get_date_range(request):
    start_raw = request.GET.get("start_date")
    end_raw = request.GET.get("end_date")

    start_date = None
    end_date = None

    if start_raw:
        # midnight start of day
        start_date = timezone.make_aware(datetime.combine(parse_date(start_raw), datetime.min.time()))
    if end_raw:
        # end of day
        end_date = timezone.make_aware(datetime.combine(parse_date(end_raw), datetime.max.time()))

    return start_date, end_date



def apply_date_filter(queryset, start_date, end_date, field="timestamp"):
    if start_date:
        queryset = queryset.filter(**{f"{field}__gte": start_date})
    if end_date:
        queryset = queryset.filter(**{f"{field}__lte": end_date})
    return queryset


# -----------------------------
# CUSTOMER REPORTS
# -----------------------------
@login_required
def customer_settlement_report(request):
    start_date, end_date = get_date_range(request)

    # Include motorcycles and agents
    customers = Customer.objects.select_related("agent").prefetch_related("motorcycles")

    # Apply date filter using a valid field (created_at or last_financial_update)
    customers = apply_date_filter(customers, start_date, end_date, "created_at")

    data = []
    for c in customers:
        for m in c.motorcycles.all():
            arrears = getattr(c, "arrears_amount", 0)
            balance_remaining = getattr(c, "balance_remaining", 0)

            # Status logic: Settled only if both arrears and balance are 0
            if arrears == 0 and balance_remaining == 0:
                status = "Settled"
            else:
                status = "Unsettled"

            data.append((
                c.full_name,
                m.vin,
                m.agent.display_name if m.agent else "",
                arrears,
                balance_remaining,
                status
            ))

    headers = ["Customer", "Motorcycle VIN", "Agent", "Arrears", "Remaining Balance", "Status"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="customer_settlement.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(data, headers, filename="customer_settlement.pdf", title="Customer Settlement Report")

    return render(request, "reports/customer_settlement.html", {
        "customers": customers,
        "report_data": data,
        "start_date": start_date,
        "end_date": end_date,
    })



@login_required
def customer_arrears_aging_report(request):
    start_date, end_date = get_date_range(request)
    customers = apply_date_filter(Customer.objects.filter(settled_flag=False), start_date, end_date, "timestamp")

    data = [(c.full_name, c.arrears_amount) for c in customers]
    headers = ["Customer", "Arrears Amount"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="customer_arrears_aging.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(data, headers, filename="customer_arrears_aging.pdf", title="Customer Arrears Aging Report")

    return render(request, "reports/customer_arrears_aging.html", {"customers": customers})


@login_required
def customer_statement_report(request, customer_id):
    start_date, end_date = get_date_range(request)
    customer = get_object_or_404(Customer, id=customer_id)
    invoices = apply_date_filter(Invoice.objects.filter(customer=customer).order_by("timestamp"), start_date, end_date)
    payments = apply_date_filter(Payment.objects.filter(customer=customer).order_by("timestamp"), start_date, end_date)

    total_invoices = sum(inv.invoice_amount for inv in invoices)
    total_payments = sum(pay.amount for pay in payments)
    closing_balance = total_invoices - total_payments

    data = []
    for inv in invoices:
        data.append((inv.timestamp.strftime("%Y-%m-%d"), "Invoice", inv.id, inv.invoice_amount))
    for pay in payments:
        data.append((pay.timestamp.strftime("%Y-%m-%d"), "Payment", pay.id, pay.amount))
    headers = ["Date", "Type", "Reference", "Amount"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename=f"customer_statement_{customer.id}.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(data, headers, filename=f"customer_statement_{customer.id}.pdf", title=f"Customer Statement - {customer.full_name}")

    return render(request, "reports/customer_statement.html", {
        "customer": customer,
        "invoices": invoices,
        "payments": payments,
        "total_invoices": total_invoices,
        "total_payments": total_payments,
        "closing_balance": closing_balance,
    })


@login_required
def customer_aggregation_report(request):
    start_date, end_date = get_date_range(request)

    # Base queryset
    customers = Customer.objects.annotate(
        total_invoices=Sum("invoices__amount_due"),   # Invoice field
        total_payments=Sum("payments__amount_paid"),  # Payment field
        total_credit=Sum("payments__credit"),         # Credit column
        arrears=F("arrears_amount"),                  # Customer field
    )

    # Apply date filter (from partial)
    customers = apply_date_filter(customers, start_date, end_date, "created_at")

    # Optional filter by customer name (from partial)
    search_name = request.GET.get("search_name")
    if search_name:
        customers = customers.filter(full_name__icontains=search_name)

    # Prepare data for export
    data = [
        (c.full_name, c.total_invoices, c.total_payments, c.total_credit, c.arrears)
        for c in customers
    ]
    headers = ["Customer", "Total Invoices", "Total Payments", "Total Credit", "Arrears"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="customer_aggregation.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(
            data, headers,
            filename="customer_aggregation.pdf",
            title="Customer Aggregation Report"
        )

    return render(
        request,
        "reports/customer_aggregation.html",
        {
            "customers": customers,
            "report_type": "customer_aggregation",  # ✅ tells partial to show search box
        }
    )




# -----------------------------
# AGENT REPORTS
# -----------------------------
@login_required
def agent_commission_report(request):
    start_date, end_date = get_date_range(request)
    payouts = apply_date_filter(CommissionPayout.objects.select_related("agent").all(), start_date, end_date, "paid_on")

    data = [(p.agent.full_name, p.amount, p.paid_on.strftime("%Y-%m-%d %H:%M")) for p in payouts]
    headers = ["Agent", "Amount", "Payout Date"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="agent_commission.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(data, headers, filename="agent_commission.pdf", title="Agent Commission Report")

    return render(request, "reports/agent_commission.html", {"commission_payouts": payouts})

@login_required
def agent_performance_report(request):
    start_date, end_date = get_date_range(request)

    agents = Agent.objects.annotate(
        sales_count=Count("invoices", distinct=True),
        customer_count=Count("customers", distinct=True),
    )

    agents = apply_date_filter(agents, start_date, end_date, "created_on")

    def get_rating(sales_count):
        if sales_count == 0:
            return "Bad"
        elif sales_count <= 2:
            return "Poor"
        elif sales_count <= 5:
            return "At least"
        elif sales_count <= 10:
            return "Reasonable"
        elif sales_count <= 20:
            return "Satisfactory"
        elif sales_count <= 30:
            return "Best"
        else:
            return "Excellent"

    # ✅ Attach rating to each agent object
    for agent in agents:
        agent.rating = get_rating(agent.sales_count)

    # Prepare data for export
    data = [
        (a.display_name, a.sales_count, a.customer_count, a.rating)
        for a in agents
    ]
    headers = ["Agent", "Sales Count", "Customer Count", "Rating"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="agent_performance.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(
            data, headers,
            filename="agent_performance.pdf",
            title="Agent Performance Report"
        )

    return render(
        request,
        "reports/agent_performance.html",
        {
            "agents": agents,
            "report_type": "agent_performance",
        }
    )



# -----------------------------
# INVOICES REPORTS
# -----------------------------
@login_required
def invoice_status_report(request):
    stats = Invoice.objects.values("status").annotate(count=Count("id"))
    data = [(s["status"], s["count"]) for s in stats]
    headers = ["Status", "Count"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="invoice_status.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(data, headers, filename="invoice_status.pdf", title="Invoice Status Report")

    return render(request, "reports/invoice_status.html", {"invoice_stats": stats})

# ----------------------------
# Invoice Aging Helpers (flat rows)
# ----------------------------
def export_invoice_aging_excel(rows, headers, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    import io
    from django.http import HttpResponse

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice Aging"

    # --- Add headers ---
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="AAAAAA", end_color="AAAAAA", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # --- Add rows with conditional formatting ---
    for row in rows:
        ws.append([str(cell) for cell in row])
        last_row = ws[ws.max_row]

        # ✅ Overdue status is the last column (index -1)
        status_cell = last_row[-1]
        if "days" in status_cell.value:  # overdue
            status_cell.fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
            status_cell.font = Font(color="9C0006", bold=True)
        else:  # On Time
            status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            status_cell.font = Font(color="006100", bold=True)

    # --- Auto-adjust column widths ---
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_length + 2

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(output.read(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response



def export_invoice_aging_pdf(rows, headers, filename, title, start_date=None, end_date=None, summary_rows=None):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    import io
    from django.http import HttpResponse

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()

    bold_style = ParagraphStyle('BoldHeader', parent=styles['Normal'], fontName='Helvetica-Bold')

    # --- Title ---
    elements.append(Paragraph(title, styles['Title']))
    if start_date and end_date:
        elements.append(Paragraph(f"Period: {start_date.strftime('%Y-%m-%d')} to {end_date.strftime('%Y-%m-%d')}", styles['Normal']))
    elements.append(Spacer(1, 12))

    # --- Summary block ---
    if summary_rows:
        for row in summary_rows:
            elements.append(Paragraph(f"{row[0]}: {row[1]}", styles['Normal']))
        elements.append(Spacer(1, 12))

    # --- Table data ---
    table_data = [[Paragraph(str(h), bold_style) for h in headers]]
    for row in rows:
        table_data.append([Paragraph(str(cell), styles['Normal']) for cell in row])

    page_width = A4[0] - doc.leftMargin - doc.rightMargin
    col_widths = [page_width / len(headers)] * len(headers)

    t = Table(table_data, colWidths=col_widths, repeatRows=1)

    style = TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.25, colors.black),
        ('BOX', (0,0), (-1,-1), 0.25, colors.black),
        ('BACKGROUND', (0,0), (-1,0), colors.grey),
    ])

    # ✅ Conditional formatting for Overdue Status column (last column)
    for r_idx, row in enumerate(rows, start=1):  # start=1 because header is row 0
        status = row[-1]
        if "days" in status:  # overdue
            style.add('BACKGROUND', (-1, r_idx), (-1, r_idx), colors.lightcoral)
            style.add('TEXTCOLOR', (-1, r_idx), (-1, r_idx), colors.red)
        else:  # On Time
            style.add('BACKGROUND', (-1, r_idx), (-1, r_idx), colors.lightgreen)
            style.add('TEXTCOLOR', (-1, r_idx), (-1, r_idx), colors.green)

    t.setStyle(style)
    elements.append(t)

    doc.build(elements)

    pdf_value = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf_value, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response



@login_required
def invoice_aging_report(request):
    start_date, end_date = get_date_range(request)

    invoices = apply_date_filter(
        Invoice.objects.exclude(status__in=["PAID", "FAILED"])
        .select_related("customer", "motorcycle", "motorcycle__agent"),
        start_date, end_date, "due_date"
    )

    today = now().date()
    data = []
    total_overdue_amount = 0
    total_invoices = 0

    for i in invoices:
        due_date = i.due_date or i.created_at.date()
        days_since_creation = (today - i.created_at.date()).days

        if today > due_date:
            days_overdue = (today - due_date).days
            overdue_display = f"{days_overdue} days"
        else:
            days_overdue = 0
            overdue_display = "On Time"

        total_invoices += 1
        if days_overdue > 0:
            total_overdue_amount += i.amount_due

        row = [
            str(i.invoice_id),
            str(i.customer.full_name),
            str(i.motorcycle.vin) if i.motorcycle else "",
            str(i.motorcycle.agent.display_name) if i.motorcycle and i.motorcycle.agent else "",
            f"{i.amount_due:.2f}",
            i.created_at.strftime("%Y-%m-%d"),
            str(days_since_creation),
            overdue_display
        ]
        data.append(row)

    headers = [
        "Invoice ID",
        "Customer",
        "Motorcycle VIN",
        "Agent",
        "Amount",
        "Date",
        "Days Since Creation",
        "Overdue Status"
    ]

    if request.GET.get("format") == "excel":
        return export_invoice_aging_excel(data, headers, filename="invoice_aging.xlsx")

    elif request.GET.get("format") == "pdf":
        summary_rows = [
            ["Total Invoices", str(total_invoices)],
            ["Total Overdue Amount", f"{total_overdue_amount:.2f}"],
        ]
        return export_invoice_aging_pdf(
            data,
            headers,
            filename="invoice_aging.pdf",
            title="Invoice Aging Report",
            start_date=start_date,
            end_date=end_date,
            summary_rows=summary_rows
        )

    # ✅ Pass summary values to template so browser shows them too
    return render(request, "reports/invoice_aging.html", {
        "invoices": invoices,
        "report_data": data,
        "headers": headers,
        "start_date": start_date,
        "end_date": end_date,
        "total_invoices": total_invoices,
        "total_overdue_amount": total_overdue_amount,
    })





# -----------------------------
# PAYMENT REPORTS
# -----------------------------
@login_required
def payment_success_report(request):
    start_date, end_date = get_date_range(request)
    payments = apply_date_filter(Payment.objects.all(), start_date, end_date)

    stats = payments.values("status").annotate(count=Count("id"))
    data = [(s["status"], s["count"]) for s in stats]
    headers = ["Status", "Count"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="payment_success.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(data, headers, filename="payment_success.pdf", title="Payment Success Report")

    return render(request, "reports/payment_success.html", {"payments": stats})


@login_required
def payment_reconciliation_report(request):
    # Get search query for VIN
    vin_query = request.GET.get("vin", "").strip()

    # Base queryset
    invoices = Invoice.objects.select_related("motorcycle").all()
    payments = Payment.objects.select_related("invoice__motorcycle").all()

    # ✅ Apply VIN search filter
    if vin_query:
        invoices = invoices.filter(motorcycle__vin__icontains=vin_query)
        payments = payments.filter(invoice__motorcycle__vin__icontains=vin_query)

    # ✅ Totals
    total_invoices = sum(inv.amount_due for inv in invoices)
    total_payments = sum(pay.amount for pay in payments)
    net_balance = sum(inv.balance_remaining for inv in invoices)

    # ✅ Prepare export data
    data = []
    for inv in invoices:
        data.append((
            inv.motorcycle.vin if inv.motorcycle else "N/A",
            f"Invoice ({inv.invoice_id})",
            inv.amount_due,
            inv.status,
            inv.created_at.strftime("%Y-%m-%d"),
            inv.balance_remaining
        ))
    for pay in payments:
        data.append((
            pay.invoice.motorcycle.vin if pay.invoice and pay.invoice.motorcycle else "N/A",
            f"Payment ({pay.id})",
            pay.amount,
            pay.status,
            pay.timestamp.strftime("%Y-%m-%d"),
            "-"
        ))

    headers = ["Reference (VIN)", "Type (ID)", "Amount", "Status", "Date", "Balance"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="payment_reconciliation.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(
            data, headers,
            filename="payment_reconciliation.pdf",
            title="Payment Reconciliation Report"
        )

    return render(
        request,
        "reports/payment_reconciliation.html",
        {
            "invoices": invoices,
            "payments": payments,
            "total_invoices": total_invoices,
            "total_payments": total_payments,
            "net_balance": net_balance,
            "vin_query": vin_query,  # ✅ pass search term back to template
        }
    )

# -----------------------------
# ACCOUNTING REPORTS
# -----------------------------
@login_required
def ledger_report(request):
    start_date, end_date = get_date_range(request)
    entries = apply_date_filter(LedgerEntry.objects.select_related("customer", "motorcycle").all(), start_date, end_date)

    data = [
        (
            e.customer.full_name if e.customer else "",
            e.motorcycle.vin if e.motorcycle else "",
            e.entry_type,
            e.amount,
            e.timestamp.strftime("%Y-%m-%d")
        )
        for e in entries
    ]
    headers = ["Customer", "Motorcycle VIN", "Entry Type", "Amount", "Date"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="ledger_report.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(data, headers, filename="ledger_report.pdf", title="Ledger Report")

    return render(request, "reports/ledger.html", {"entries": entries})


@login_required
def credit_vs_arrears_report(request):
    # ✅ Filters
    search_customer = request.GET.get("customer", "").strip()
    search_vin = request.GET.get("vin", "").strip()

    # Base queryset
    accounts = CustomerAccount.objects.select_related("customer", "motorcycle").all()

    # Apply filters
    if search_customer:
        accounts = accounts.filter(customer__full_name__icontains=search_customer)
    if search_vin:
        accounts = accounts.filter(motorcycle__vin__icontains=search_vin)

    # ✅ Totals
    total_credit = accounts.aggregate(total_credit=Sum("credit_balance"))["total_credit"] or 0
    total_arrears = accounts.aggregate(total_arrears=Sum("arrears_balance"))["total_arrears"] or 0
    net_balance = total_credit - total_arrears

    # ✅ Attach net_balance to each account
    for acc in accounts:
        acc.net_balance = (acc.credit_balance or 0) - (acc.arrears_balance or 0)

    # ✅ Prepare export data
    data = []
    for acc in accounts:
        data.append((
            acc.customer.full_name,
            acc.motorcycle.vin if acc.motorcycle else "N/A",
            acc.credit_balance,
            acc.arrears_balance,
            acc.net_balance
        ))

    headers = ["Customer", "Motorcycle VIN", "Credit Balance", "Arrears Balance", "Net Balance"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="credit_vs_arrears.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(
            data, headers,
            filename="credit_vs_arrears.pdf",
            title="Credit vs Arrears Report"
        )

    return render(
        request,
        "reports/credit_vs_arrears.html",
        {
            "accounts": accounts,
            "total_credit": total_credit,
            "total_arrears": total_arrears,
            "net_balance": net_balance,
            "search_customer": search_customer,
            "search_vin": search_vin,
        }
    )


# -----------------------------
# WAREHOUSE REPORTS
# -----------------------------
@login_required
def warehouse_stock_report(request):
    start_date, end_date = get_date_range(request)
    warehouses = Warehouse.objects.annotate(count=Count("motorcycles"))
    warehouses = apply_date_filter(warehouses, start_date, end_date, "timestamp")

    data = [(w.physical_location, w.count) for w in warehouses]
    headers = ["Warehouse", "Stock Count"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="warehouse_stock.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(data, headers, filename="warehouse_stock.pdf", title="Warehouse Stock Report")

    return render(request, "reports/warehouse_stock.html", {"warehouse_stats": warehouses})


# -----------------------------
# MOTORCYCLE REPORTS
# -----------------------------
@login_required
def motorcycle_status_report(request):
    start_date, end_date = get_date_range(request)
    motorcycles = apply_date_filter(Motorcycle.objects.all(), start_date, end_date)
    stats = motorcycles.values("status").annotate(count=Count("id"))

    data = [(s["status"], s["count"]) for s in stats]
    headers = ["Status", "Count"]

    if request.GET.get("format") == "excel":
        return export_to_excel(data, headers, filename="motorcycle_status.xlsx")
    elif request.GET.get("format") == "pdf":
        return export_to_pdf(data, headers, filename="motorcycle_status.pdf", title="Motorcycle Status Report")

    return render(request, "reports/motorcycle_status.html", {"motorcycle_stats": stats})



@login_required
def motorcycle_movement_report(request):
    start_date, end_date = get_date_range(request)
    vin = request.GET.get("vin", "").strip()

    motorcycle = None
    events = []

    if vin:
        try:
            motorcycle = Motorcycle.objects.get(vin__iexact=vin)
            history = apply_date_filter(
                motorcycle.history.all().order_by("timestamp"),
                start_date, end_date
            )

            for h in history:
                user_info = h.performed_by.get_full_name() if h.performed_by else "System"
                events.append({
                    "date": h.timestamp,
                    "event": h.action,
                    "details": f"{h.notes or ''} (by {user_info})"
                })

            # Sort chronologically
            events.sort(key=lambda e: e["date"])

            headers = ["Date", "Event", "Details"]
            # ✅ Short date format for rows
            data = [(e["date"].strftime("%Y-%m-%d"), e["event"], e["details"]) for e in events]

            # ✅ Export branches with short date formatting
            if request.GET.get("format") == "excel":
                return export_movement_excel(
                    data,
                    headers,
                    filename=f"motorcycle_movement_{vin}_{start_date.strftime('%Y-%m-%d')}_{end_date.strftime('%Y-%m-%d')}.xlsx"
                )

            elif request.GET.get("format") == "pdf":
                return export_movement_pdf(
                    data,
                    headers,
                    filename=f"motorcycle_movement_{vin}_{start_date.strftime('%Y-%m-%d')}_{end_date.strftime('%Y-%m-%d')}.pdf",
                    title=f"Motorcycle Movement Report - {motorcycle.vin} ({start_date.strftime('%Y-%m-%d')} to {end_date.strftime('%Y-%m-%d')})",
                    start_date=start_date.strftime('%Y-%m-%d'),
                    end_date=end_date.strftime('%Y-%m-%d')
                )

        except Motorcycle.DoesNotExist:
            motorcycle = None

    return render(request, "reports/motorcycle_movement.html", {
        "motorcycle": motorcycle,
        "events": events,
        "vin": vin,
        "start_date": start_date,
        "end_date": end_date,
    })






# ----------------------------
# Motorcycle Statement Helpers (summary + transactions)
# ----------------------------
def export_motorcycle_statement_excel(summary_rows, transaction_headers, transaction_rows, filename):
    """
    summary_rows: list of 2-element lists, e.g. [["Motorcycle VIN", "MC 123 AAA"], ["Customer", "Mwamba"], ...]
    transaction_headers: list of strings
    transaction_rows: list of lists, each matching transaction_headers length
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    import io
    from django.http import HttpResponse

    wb = Workbook()
    ws = wb.active
    ws.title = "Motorcycle Statement"

    # --- Summary block ---
    for row in summary_rows:
        ws.append([str(cell) for cell in row])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

    ws.append([])  # blank line

    # --- Transaction header ---
    ws.append(transaction_headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)  # bold header text
        cell.fill = PatternFill(start_color="AAAAAA", end_color="AAAAAA", fill_type="solid")

    # --- Transaction rows ---
    for row in transaction_rows:
        ws.append([str(cell) for cell in row])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(output.read(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response



def export_motorcycle_statement_pdf(summary_rows, transaction_headers, transaction_rows, filename, title):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    import io
    from django.http import HttpResponse

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()

    # Bold style for headers
    bold_style = ParagraphStyle(
        'BoldHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold'
    )

    elements.append(Paragraph(title, styles['Title']))
    elements.append(Spacer(1, 12))

    table_data = []

    # --- Summary block ---
    for row in summary_rows:
        table_data.append([Paragraph(str(row[0]), styles['Normal']),
                           Paragraph(str(row[1]), styles['Normal'])])
    table_data.append([])  # blank line

    # --- Transaction header ---
    table_data.append([Paragraph(str(h), bold_style) for h in transaction_headers])

    # --- Transaction rows ---
    for row in transaction_rows:
        table_data.append([Paragraph(str(cell), styles['Normal']) for cell in row])

    # Column widths
    page_width = A4[0] - doc.leftMargin - doc.rightMargin
    col_widths = [page_width / len(transaction_headers)] * len(transaction_headers)

    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    style = TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('ALIGN', (3,2), (3,-1), 'RIGHT'),   # Amount column
        ('ALIGN', (4,2), (4,-1), 'RIGHT'),   # Commission column
        ('ALIGN', (5,2), (5,-1), 'RIGHT'),   # Balance column
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.25, colors.black),
        ('BOX', (0,0), (-1,-1), 0.25, colors.black),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BACKGROUND', (0,0), (-1,0), colors.grey),
    ])
    t.setStyle(style)
    elements.append(t)

    doc.build(elements)

    pdf_value = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf_value, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response



# --- Main view ---
@login_required
def motorcycle_statement_report(request):
    start_date, end_date = get_date_range(request)
    vin = request.GET.get("vin", "").strip()

    motorcycle = None
    combined_entries = []
    total_invoices = 0
    total_payments = 0
    total_commission = 0
    opening_balance = 0
    closing_balance = 0

    if vin:
        try:
            motorcycle = Motorcycle.objects.select_related("customer_account", "agent").get(vin__iexact=vin)

            if motorcycle.customer_account:
                opening_balance = getattr(motorcycle.customer_account, "opening_balance", 0)

            invoices = apply_date_filter(
                Invoice.objects.filter(motorcycle=motorcycle).order_by("created_at"),
                start_date, end_date, field="created_at"
            )
            payments = apply_date_filter(
                Payment.objects.filter(motorcycle=motorcycle).order_by("timestamp"),
                start_date, end_date, field="timestamp"
            )
            ledger_entries = apply_date_filter(
                LedgerEntry.objects.filter(motorcycle=motorcycle).order_by("created_at"),
                start_date, end_date, field="created_at"
            )

            for inv in invoices:
                combined_entries.append({
                    "date": inv.created_at,
                    "type": "Invoice",
                    "reference": inv.invoice_id,
                    "amount": inv.amount_due,
                    "commission": getattr(inv, "agent_commission", 0),
                })
                total_invoices += inv.amount_due
                total_commission += getattr(inv, "agent_commission", 0)

            for pay in payments:
                combined_entries.append({
                    "date": pay.timestamp,
                    "type": "Payment",
                    "reference": pay.transaction_id,
                    "amount": pay.amount_paid,
                    "commission": None,
                })
                total_payments += pay.amount_paid

            for entry in ledger_entries:
                combined_entries.append({
                    "date": entry.created_at,
                    "type": "Ledger Entry",
                    "reference": entry.entry_type,
                    "amount": entry.amount,
                    "commission": None,
                })

            combined_entries.sort(key=lambda e: e["date"])

            running_balance = opening_balance
            for e in combined_entries:
                if e["type"] == "Payment":
                    running_balance -= e["amount"]
                e["running_balance"] = running_balance

            closing_balance = opening_balance - total_payments

            # --- Summary rows ---
            summary_rows = [
                ["Motorcycle VIN", motorcycle.vin],
                ["Customer", getattr(motorcycle.customer_account, "full_name", "")],
                ["Agent", motorcycle.agent.display_name if motorcycle.agent else ""],
                ["Period", f"{start_date.strftime('%Y-%m-%d')} to {end_date.strftime('%Y-%m-%d')}" if start_date and end_date else ""],
                ["Opening Balance", number_format(opening_balance, 2)],
                ["Total Invoices", number_format(total_invoices, 2)],
                ["Total Payments", number_format(total_payments, 2)],
                ["Total Commission", number_format(total_commission, 2)],
                ["Closing Balance", number_format(closing_balance, 2)],
            ]

            # --- Transaction headers ---
            transaction_headers = ["Date", "Type", "Reference", "Amount (Tsh)", "Agent Commission (Tsh)", "Running Balance (Tsh)"]

            # --- Transaction rows ---
            transaction_rows = [
                [
                    e["date"].strftime("%Y-%m-%d"),
                    e["type"],
                    e["reference"],
                    number_format(e["amount"], decimal_pos=2, use_l10n=True),
                    number_format(e["commission"], decimal_pos=2, use_l10n=True) if e["commission"] else "",
                    number_format(e["running_balance"], decimal_pos=2, use_l10n=True),
                ]
                for e in combined_entries
            ]

            # --- Export ---
            if request.GET.get("format") == "excel":
                return export_motorcycle_statement_excel(
                    summary_rows,
                    transaction_headers,
                    transaction_rows,
                    filename=f"motorcycle_statement_{vin}.xlsx"
                )

            elif request.GET.get("format") == "pdf":
                title = f"Motorcycle Statement - {motorcycle.vin}"
                if start_date and end_date:
                    title += f" ({start_date.strftime('%Y-%m-%d')} to {end_date.strftime('%Y-%m-%d')})"
                return export_motorcycle_statement_pdf(
                    summary_rows,
                    transaction_headers,
                    transaction_rows,
                    filename=f"motorcycle_statement_{vin}.pdf",
                    title=title
                )

        except Motorcycle.DoesNotExist:
            motorcycle = None

    return render(request, "reports/motorcycle_statement.html", {
        "motorcycle": motorcycle,
        "combined_entries": combined_entries,
        "vin": vin,
        "start_date": start_date,
        "end_date": end_date,
        "opening_balance": opening_balance,
        "total_invoices": total_invoices,
        "total_payments": total_payments,
        "total_commission": total_commission,
        "closing_balance": closing_balance,
        "customer_name": getattr(motorcycle.customer_account, "full_name", None) if motorcycle and motorcycle.customer_account else None,
        "agent_name": motorcycle.agent.display_name if motorcycle and motorcycle.agent else None,
    })

def style_excel(ws):
    # Bold header with background color
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Auto column width
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = max_length + 2



@login_required
def outstanding_report(request):
    start_date, end_date = get_date_range(request)
    customer = request.GET.get("customer", "").strip()
    vin = request.GET.get("vin", "").strip()

    # ✅ Fetch all invoices except PAID and FAILED
    invoices = Invoice.objects.exclude(status__in=["PAID", "FAILED"]) \
        .select_related("customer", "motorcycle", "motorcycle__agent")

    # ✅ Apply date filter only if provided
    if start_date and end_date:
        invoices = invoices.filter(created_at__range=(start_date, end_date))

    # ✅ Apply optional filters
    if customer:
        invoices = invoices.filter(customer__full_name__icontains=customer)
    if vin:
        invoices = invoices.filter(motorcycle__vin__iexact=vin)

    report_data = []
    total_outstanding = total_arrears = total_commission = 0
    customers, motorcycles = set(), set()

    for inv in invoices:
        # ✅ Correct arrears logic
        arrears = 0
        if inv.status == "PARTIAL":
            arrears = inv.balance_remaining
        elif inv.status in ["OVERDUE", "DEFAULTED"]:
            arrears = inv.balance_remaining
        # PENDING invoices → arrears = 0

        agent = inv.motorcycle.agent if inv.motorcycle and inv.motorcycle.agent else None
        commission = float(inv.agent_commission) if inv.agent_commission else 0

        total_outstanding += float(inv.amount_due)
        total_arrears += float(arrears)
        total_commission += commission
        customers.add(inv.customer.id)
        if inv.motorcycle:
            motorcycles.add(inv.motorcycle.id)

        report_data.append({
            "invoice_id": inv.invoice_id,
            "customer": inv.customer.full_name,
            "vin": inv.motorcycle.vin if inv.motorcycle else "",
            "agent": agent.display_name if agent else "",
            "amount_due": f"{inv.amount_due:.2f}",
            "status": inv.status,
            "arrears": f"{arrears:.2f}",
            "commission": f"{commission:.2f}",
        })

    # ✅ Rating logic based on arrears vs outstanding
    rating_percentage = (1 - (total_arrears / total_outstanding)) * 100 if total_outstanding > 0 else 0
    if rating_percentage >= 90:
        rating = "Excellent"
    elif rating_percentage >= 75:
        rating = "Good"
    elif rating_percentage >= 50:
        rating = "Satisfactory"
    else:
        rating = "Poor"

    summary = {
        "total_outstanding": total_outstanding,
        "total_customers": len(customers),
        "total_motorcycles": len(motorcycles),
        "total_arrears": total_arrears,
        "total_commission": total_commission,
        "rating": rating,
        "rating_percentage": f"{rating_percentage:.2f}%",
    }

    return render(request, "reports/outstanding.html", {
        "report_data": report_data,
        "summary": summary,
        "start_date": start_date,
        "end_date": end_date,
    })


@login_required
def collection_report(request):
    start_date, end_date = get_date_range(request)
    customer = request.GET.get("customer", "").strip()
    vin = request.GET.get("vin", "").strip()

    # ✅ Fetch all payments, apply date filter only if provided
    payments = Payment.objects.select_related("customer", "motorcycle", "motorcycle__agent", "invoice")
    if start_date and end_date:
        payments = payments.filter(timestamp__range=(start_date, end_date))

    # ✅ Apply optional filters
    if customer:
        payments = payments.filter(customer__full_name__icontains=customer)
    if vin:
        payments = payments.filter(motorcycle__vin__iexact=vin)

    report_data = []
    total_collected = total_credit = total_commission = 0
    customers, motorcycles = set(), set()

    for pay in payments:
        invoice = pay.invoice
        agent = pay.motorcycle.agent if pay.motorcycle and pay.motorcycle.agent else None

        # ✅ Credit logic
        credit = 0
        if invoice and pay.amount_paid > invoice.amount_due:
            credit = pay.amount_paid - invoice.amount_due

        # ✅ Commission logic
        commission = (pay.amount_paid * agent.commission_rate / 100) if agent and hasattr(agent, "commission_rate") else 0

        # ✅ Totals
        total_collected += float(pay.amount_paid)
        total_credit += float(credit)
        total_commission += float(commission)
        customers.add(pay.customer.id)
        if pay.motorcycle:
            motorcycles.add(pay.motorcycle.id)

        report_data.append({
            "transaction_id": pay.transaction_id,
            "customer": pay.customer.full_name,
            "vin": pay.motorcycle.vin if pay.motorcycle else "",
            "agent": agent.display_name if agent else "",
            "invoice_id": invoice.invoice_id if invoice else "",
            "amount_paid": f"{pay.amount_paid:.2f}",
            "credit": f"{credit:.2f}",
            "commission": f"{commission:.2f}",
            "timestamp": pay.timestamp.strftime("%Y-%m-%d"),
            "method": pay.method,
        })

    # ✅ Collection rate logic
    total_invoiced = Invoice.objects.filter(created_at__range=(start_date, end_date)) \
        .aggregate(Sum("amount_due"))["amount_due__sum"] or 0
    collection_rate = (total_collected / total_invoiced * 100) if total_invoiced > 0 else 0

    if collection_rate >= 90:
        rating = "Excellent"
    elif collection_rate >= 75:
        rating = "Good"
    elif collection_rate >= 50:
        rating = "Satisfactory"
    else:
        rating = "Poor"

    summary = {
        "total_collected": total_collected,
        "total_credit": total_credit,
        "total_commission": total_commission,
        "total_customers": len(customers),
        "total_motorcycles": len(motorcycles),
        "collection_rate": f"{collection_rate:.2f}%",
        "rating": rating,
    }

    return render(request, "reports/collection.html", {
        "report_data": report_data,
        "summary": summary,
        "start_date": start_date,
        "end_date": end_date,
    })
