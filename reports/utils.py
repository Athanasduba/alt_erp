import io
import os
from django.http import HttpResponse
from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet

# -------------------------
# Generic Excel Export
# -------------------------
def export_to_excel(data, headers, filename="report.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in data:
        ws.append(row)

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


# -------------------------
# Generic PDF Export
# -------------------------
def export_to_pdf(data, headers, filename="report.pdf", title="ERP Report"):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()

    # Logo
    logo_path = os.path.join(settings.BASE_DIR, "static/img/logo.png")
    if os.path.exists(logo_path):
        elements.append(Paragraph(f'<img src="{logo_path}" width="100" height="50"/>', styles['Normal']))

    # Title
    elements.append(Paragraph(title, styles['Title']))
    elements.append(Spacer(1, 12))

    # Table
    table_data = [headers] + data
    t = Table(table_data, repeatRows=1)
    style = TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.grey),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.25, colors.black),
        ('BOX', (0,0), (-1,-1), 0.25, colors.black),
    ])
    t.setStyle(style)
    elements.append(t)

    doc.build(elements)
    pdf = buffer.getvalue()
    buffer.close()

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.write(pdf)
    return response


# -------------------------
# Motorcycle Movement Excel Export
# -------------------------
def export_movement_excel(rows, headers, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Motorcycle Movement"

    # Headers
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="AAAAAA", end_color="AAAAAA", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    # Rows with conditional formatting
    for row in rows:
        ws.append([str(cell) for cell in row])
        event_cell = ws[ws.max_row][1]  # Event column
        if event_cell.value == "Created":
            event_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            event_cell.font = Font(color="006100", bold=True)
        elif event_cell.value == "Warehouse":
            event_cell.fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
            event_cell.font = Font(color="1F4E78", bold=True)
        elif event_cell.value == "Movement":
            event_cell.fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
            event_cell.font = Font(color="9C6500", bold=True)
        elif event_cell.value == "Invoice":
            event_cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            event_cell.font = Font(color="1F4E78", bold=True)
        elif event_cell.value == "Payment":
            event_cell.fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
            event_cell.font = Font(color="9C0006", bold=True)

    # Auto column widths
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_length + 2

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(output.read(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# -------------------------
# Motorcycle Movement PDF Export
# -------------------------
def export_movement_pdf(rows, headers, filename, title, start_date=None, end_date=None):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()

    # Title
    elements.append(Paragraph(title, styles['Title']))
    if start_date and end_date:
        elements.append(Paragraph(f"Period: {start_date} to {end_date}", styles['Normal']))
    elements.append(Spacer(1, 12))

    # Table
    table_data = [headers] + rows
    t = Table(table_data, repeatRows=1)
    style = TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.grey),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.25, colors.black),
        ('BOX', (0,0), (-1,-1), 0.25, colors.black),
    ])

    # Conditional formatting for Event column (index 1)
    for r_idx, row in enumerate(rows, start=1):
        event = row[1]
        if event == "Created":
            style.add('BACKGROUND', (1, r_idx), (1, r_idx), colors.lightgreen)
            style.add('TEXTCOLOR', (1, r_idx), (1, r_idx), colors.green)
        elif event == "Warehouse":
            style.add('BACKGROUND', (1, r_idx), (1, r_idx), colors.lightblue)
            style.add('TEXTCOLOR', (1, r_idx), (1, r_idx), colors.blue)
        elif event == "Movement":
            style.add('BACKGROUND', (1, r_idx), (1, r_idx), colors.lightyellow)
            style.add('TEXTCOLOR', (1, r_idx), (1, r_idx), colors.brown)
        elif event == "Invoice":
            style.add('BACKGROUND', (1, r_idx), (1, r_idx), colors.whitesmoke)
            style.add('TEXTCOLOR', (1, r_idx), (1, r_idx), colors.darkblue)
        elif event == "Payment":
            style.add('BACKGROUND', (1, r_idx), (1, r_idx), colors.lightcoral)
            style.add('TEXTCOLOR', (1, r_idx), (1, r_idx), colors.red)

    t.setStyle(style)
    elements.append(t)

    doc.build(elements)
    pdf = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
